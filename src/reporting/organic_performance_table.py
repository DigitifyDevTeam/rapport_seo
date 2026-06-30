"""Render the styled organic performance comparison table on a slide."""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Pt

from src.transform.organic_performance import OrganicPerformanceSlide

# Match scripts/build_template.py KPI overview palette (teal + violet + slate).
PRIMARY = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
CARD_SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BORDER = RGBColor(0xE2, 0xE8, 0xF0)
CURRENT_TINT = RGBColor(0xCC, 0xFB, 0xF1)
PREVIOUS_TINT = RGBColor(0xE9, 0xD5, 0xFF)


def _style_cell(
    cell,
    *,
    fill: RGBColor,
    bold: bool = False,
    size: int = 11,
    align=PP_ALIGN.CENTER,
    color: RGBColor | None = None,
) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    text_color = color or PRIMARY
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    for paragraph in tf.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = text_color
            run.font.name = "Segoe UI"


def _set_cell_text(cell, text: str, **style) -> None:
    cell.text = str(text)
    _style_cell(cell, **style)


def _apply_cell_borders(cell, *, color_hex: str = "E2E8F0",
                        width_emu: str = "12700") -> None:
    """Table cells have no ``.line`` API; borders are set on the underlying XML."""
    tc_pr = cell._tc.get_or_add_tcPr()  # noqa: SLF001
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        tag = parse_xml(
            f'<a:{edge} {nsdecls("a")} w="{width_emu}" cap="flat" cmpd="sng" '
            f'algn="ctr"><a:solidFill><a:srgbClr val="{color_hex}"/>'
            f"</a:solidFill></a:{edge}>"
        )
        tc_pr.append(tag)


def add_organic_performance_table(slide, left, top, width, height,
                                   payload: OrganicPerformanceSlide) -> None:
    """Draw the 3-column GA4 organic comparison table."""
    rows = 1 + len(payload.rows)
    cols = 3
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    col_widths = (int(width * 0.38), int(width * 0.31), int(width * 0.31))
    for idx, col_w in enumerate(col_widths):
        table.columns[idx].width = col_w

    _set_cell_text(table.cell(0, 0), "", fill=CARD_SURFACE, bold=True)
    _set_cell_text(
        table.cell(0, 1),
        payload.current_range,
        fill=CURRENT_TINT,
        bold=True,
        size=10,
        color=PRIMARY,
    )
    _set_cell_text(
        table.cell(0, 2),
        payload.previous_range,
        fill=PREVIOUS_TINT,
        bold=True,
        size=10,
        color=PRIMARY,
    )

    for row_idx, (label, current, previous) in enumerate(payload.rows, start=1):
        _set_cell_text(
            table.cell(row_idx, 0),
            label,
            fill=CARD_SURFACE,
            bold=True,
            align=PP_ALIGN.LEFT,
            color=PRIMARY,
        )
        _set_cell_text(
            table.cell(row_idx, 1),
            current,
            fill=CARD_SURFACE,
            color=PRIMARY,
        )
        _set_cell_text(
            table.cell(row_idx, 2),
            previous,
            fill=CARD_SURFACE,
            color=MUTED,
        )

    border_hex = str(CARD_BORDER)
    for row_idx in range(rows):
        for col_idx in range(cols):
            cell = table.cell(row_idx, col_idx)
            try:
                from pptx.enum.text import MSO_ANCHOR
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            except (ImportError, AttributeError):
                pass
            _apply_cell_borders(cell, color_hex=border_hex)
